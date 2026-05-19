import gradio as gr
import tempfile
import os
from datetime import datetime
from docx import Document
import docx.shared
from ai_service import analyze_submission, generate_report
from config import CATEGORIES, MODELS

def process_submission(content: str, model_name: str, api_key: str, filename: str = ""):
    if not content.strip():
        return "请输入投稿内容", None, None, "❌ 请输入投稿内容", None, None
    
    if not api_key.strip():
        return "请输入API Key", None, None, "❌ 请输入API Key", None, None
    
    try:
        analysis = analyze_submission(content, model_name, api_key)
        report = generate_report(content, analysis)
        txt_file = export_txt(report, filename)
        word_file = export_word(report, filename)
        return report, report, analysis, "✅ 分析完成", txt_file, word_file
    except Exception as e:
        return f"分析失败：{str(e)}", None, None, f"❌ 分析失败：{str(e)}", None, None

def export_txt(report: str, filename: str = ""):
    if not report:
        return None
    name = filename.strip() if filename.strip() else f"投稿分析报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    if not name.endswith('.txt'):
        name += '.txt'
    filepath = os.path.join(tempfile.gettempdir(), name)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(report)
    return filepath

def export_word(report: str, filename: str = ""):
    if not report:
        return None
    name = filename.strip() if filename.strip() else f"投稿分析报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    if not name.endswith('.docx'):
        name += '.docx'
    filepath = os.path.join(tempfile.gettempdir(), name)
    
    doc = Document()
    
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '微软雅黑')
    
    lines = report.split("\n")
    for line in lines:
        if line.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(line[2:])
            run.bold = True
            run.font.size = docx.shared.Pt(18)
        elif line.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(line[3:])
            run.bold = True
            run.font.size = docx.shared.Pt(14)
        elif line.startswith("- **"):
            parts = line[2:].split("**：")
            if len(parts) == 2:
                p = doc.add_paragraph()
                label = parts[0].replace("**", "").strip()
                run1 = p.add_run(label + "：")
                run1.bold = True
                p.add_run(parts[1])
            else:
                doc.add_paragraph(line)
        elif line.startswith("---"):
            doc.add_paragraph("─" * 30)
        elif line.strip():
            doc.add_paragraph(line)
    
    doc.save(filepath)
    return filepath

with gr.Blocks(title="AI辅助投稿分析工具", theme=gr.themes.Soft()) as demo:
    gr.Markdown("# AI辅助投稿分析工具")
    gr.Markdown("快速分析社区投稿，自动分类、摘要、评估热度")
    
    with gr.Row():
        with gr.Column(scale=2):
            model_select = gr.Dropdown(
                label="选择模型",
                choices=list(MODELS.keys()),
                value="DeepSeek"
            )
            api_input = gr.Textbox(
                label="API Key",
                type="password",
                placeholder="输入对应模型的API Key"
            )
            content_input = gr.Textbox(
                label="投稿内容",
                lines=10,
                placeholder="粘贴投稿内容（支持微信/小红书/抖音等平台内容）..."
            )
            analyze_btn = gr.Button("开始分析", variant="primary")
            status_output = gr.Textbox(label="状态", interactive=False)
        
        with gr.Column(scale=3):
            report_output = gr.Markdown(label="分析报告")
            
            filename_input = gr.Textbox(
                label="文件名",
                placeholder="自定义文件名（留空自动生成）"
            )
            with gr.Row():
                txt_btn = gr.DownloadButton("下载TXT")
                word_btn = gr.DownloadButton("下载Word")
    
    hidden_report = gr.Textbox(visible=False)
    hidden_analysis = gr.JSON(visible=False)
    
    analyze_btn.click(
        fn=process_submission,
        inputs=[content_input, model_select, api_input, filename_input],
        outputs=[report_output, hidden_report, hidden_analysis, status_output, txt_btn, word_btn],
        show_progress="full"
    )
    
    gr.Markdown("""
    ---
    **使用说明**：
    1. 选择AI模型并输入对应的API Key
    2. 粘贴投稿内容到文本框
    3. 点击"开始分析"获取报告
    4. 可导出为TXT或Word格式
    
    **API Key获取地址**：
    - DeepSeek: https://platform.deepseek.com/
    - 通义千问: https://dashscope.console.aliyun.com/
    - 智谱GLM: https://open.bigmodel.cn/
    - Kimi: https://platform.moonshot.cn/
    """)

if __name__ == "__main__":
    demo.launch()
